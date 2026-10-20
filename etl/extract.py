"""EXTRACT: baca file mentah (XLSX, CSV, JSON) lalu simpan apa adanya (teks) ke kehutanan_staging.

Tidak ada pembersihan di sini. Struktur tabel sumber hanya "diratakan" (satu baris = satu provinsi
per tahun) supaya bisa disimpan di tabel, tetapi nilai tetap teks asli.
"""
import csv
import json
import re

import openpyxl

from config import DB_STAGING, RAW_DIR, catat_etl, koneksi, run_id, sekarang


def _teks(v):
    return None if v is None else str(v).strip()


def extract_d1():
    rows = []
    for path in sorted((RAW_DIR / "xlsx").glob("D1_kawasan_hutan_*.xlsx")):
        tahun = int(re.search(r"(\d{4})", path.stem).group(1))
        ws = openpyxl.load_workbook(path, data_only=True).active
        semua = list(ws.iter_rows(values_only=True))
        header = [str(h or "").lower() for h in semua[0]]

        def kolom(kata):
            return next(i for i, h in enumerate(header) if kata in h)

        idx = {
            "sk": next((i for i, h in enumerate(header) if "tahun sk" in h), None),
            "hl": kolom("lindung"),
            "ks": kolom("suaka"),
            "hpt": kolom("terbatas"),
            "hp": kolom("tetap"),
            "hpk": kolom("dikonversi"),
            "jml": kolom("jumlah"),
        }
        for n, r in enumerate(semua[1:], start=2):
            if not r[0] or str(r[0]).strip().lower().startswith(("keterangan", "<sup", "1", "2")):
                continue
            rows.append((
                path.name, tahun, n, _teks(r[0]),
                _teks(r[idx["sk"]]) if idx["sk"] is not None else None,
                _teks(r[idx["hl"]]), _teks(r[idx["ks"]]), _teks(r[idx["hpt"]]),
                _teks(r[idx["hp"]]), _teks(r[idx["hpk"]]), _teks(r[idx["jml"]]),
            ))
    return rows


def extract_d2():
    path = RAW_DIR / "csv" / "D2_tutupan_lahan_2014_2023.csv"
    semua = list(csv.reader(open(path, encoding="utf-8-sig")))
    rows = []
    blok, tahun_kolom = None, {}
    for n, r in enumerate(semua, start=1):
        if r and r[0].strip() == "No":
            # baris judul tahun: 2014 di kolom 2, 2015 di kolom 7, dst (5 kolom per tahun)
            tahun_kolom = {i: int(v) for i, v in enumerate(r) if v.strip().isdigit()}
            blok = "2014-2022" if 2014 in tahun_kolom.values() else "2023"
            continue
        if not tahun_kolom or len(r) < 3:
            continue
        nama = r[1].strip() or r[0].strip()
        if not nama or nama.startswith("(") or not any(c.strip() for c in r[2:]):
            continue
        if nama.lower().startswith(("keterangan", "sumber")):
            continue
        for i, tahun in tahun_kolom.items():
            if not r[i].strip():
                continue
            rows.append((path.name, blok, n, nama, tahun, r[i], r[i + 1], r[i + 2], r[i + 3], r[i + 4]))
    return rows


def extract_d3():
    path = RAW_DIR / "xlsx" / "D3_karhutla_provinsi_2018_2026.xlsx"
    ws = openpyxl.load_workbook(path, data_only=True).active
    semua = list(ws.iter_rows(values_only=True))
    tahun_kolom = {i: int(v) for i, v in enumerate(semua[0]) if str(v).strip().isdigit()}
    rows = []
    for n, r in enumerate(semua[1:], start=2):
        if not r[1]:
            continue
        for i, tahun in tahun_kolom.items():
            rows.append((path.name, n, _teks(r[0]), _teks(r[1]), tahun, _teks(r[i])))
    return rows


def baris_dari_json_bps(data, nama_sumber):
    var = data["var"][0]["val"]
    turvar = data["turvar"][0]["val"]
    turtahun = data["turtahun"][0]["val"]
    rows = []
    for th in data["tahun"]:
        for v in data["vervar"]:
            kunci = f"{v['val']}{var}{turvar}{th['val']}{turtahun}"
            if kunci in data["datacontent"]:
                rows.append((nama_sumber, kunci, str(v["val"]), v["label"], int(th["label"]),
                             str(data["datacontent"][kunci])))
    return rows


def extract_d4():
    rows = []
    for path in sorted((RAW_DIR / "json").glob("D4_produksi_kayu_bulat_*.json")):
        rows += baris_dari_json_bps(json.loads(path.read_text(encoding="utf-8")), path.name)
    return rows


def extract_d5():
    path = RAW_DIR / "csv" / "D5_deforestasi_netto_2013_2022.csv"
    semua = list(csv.reader(open(path, encoding="utf-8-sig")))
    baris_periode = next(r for r in semua if r and r[0].strip() == "Provinsi")
    periode_kolom = {i: v.strip() for i, v in enumerate(baris_periode) if re.match(r"\s*\d{4}-\d{4}", v)}
    rows = []
    for n, r in enumerate(semua, start=1):
        nama = r[0].strip() if r else ""
        if not nama or nama in ("Provinsi",) or nama.lower().startswith(("ket", "sumber", "angka deforestasi")):
            continue
        if len(r) < 4 or not r[1].strip():
            continue
        for i, periode in periode_kolom.items():
            rows.append((path.name, n, nama, periode, r[i], r[i + 1], r[i + 2]))
    return rows


TABEL = {
    "D1": ("stg_d1_kawasan_hutan", extract_d1,
           "sumber_file, tahun_data, baris_ke, provinsi_raw, tahun_sk_raw, hutan_lindung_raw, "
           "konservasi_raw, hpt_raw, hp_raw, hpk_raw, jumlah_raw"),
    "D2": ("stg_d2_tutupan_lahan", extract_d2,
           "sumber_file, blok, baris_ke, provinsi_raw, tahun_data, hutan_raw, hutan_pct_raw, "
           "non_hutan_raw, non_hutan_pct_raw, jumlah_raw"),
    "D3": ("stg_d3_karhutla", extract_d3,
           "sumber_file, baris_ke, kode_wilayah_raw, provinsi_raw, tahun_data, luas_raw"),
    "D4": ("stg_d4_produksi_kayu", extract_d4,
           "sumber_file, datacontent_key, vervar_val, provinsi_raw, tahun_data, nilai_raw"),
    "D5": ("stg_d5_deforestasi", extract_d5,
           "sumber_file, baris_ke, provinsi_raw, periode_raw, dalam_kawasan_raw, luar_kawasan_raw, total_raw"),
}


def main(rid=None):
    rid = rid or run_id()
    conn = koneksi(DB_STAGING)
    waktu = sekarang()
    try:
        for dataset, (tabel, fungsi, kolom) in TABEL.items():
            rows = [r + (waktu,) for r in fungsi()]
            isian = ", ".join(["%s"] * (kolom.count(",") + 2))
            with conn.cursor() as cur:
                cur.execute(f"TRUNCATE TABLE {tabel}")
                cur.executemany(f"INSERT INTO {tabel} ({kolom}, extracted_at) VALUES ({isian})", rows)
            catat_etl(conn, rid, "extract", dataset, len(rows), len(rows), 0, "OK",
                      f"{len(rows)} baris teks mentah -> {DB_STAGING}.{tabel}")
            print(f"[extract] {dataset}: {len(rows)} baris -> {tabel}")
        conn.commit()
    finally:
        conn.close()
    return rid


if __name__ == "__main__":
    main()
